import fitz  # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument
from typing import Optional
from core.logging import get_logger

logger = get_logger(__name__)


def parse_pdf(file_path: str) -> str:
    """
    Extract text from PDF file using PyMuPDF with pdfplumber fallback
    """
    try:
        # Try PyMuPDF first (faster)
        text = _parse_pdf_pymupdf(file_path)
        if text and len(text.strip()) > 0:
            return text
    except Exception as e:
        logger.warning(f"PyMuPDF failed, trying pdfplumber: {e}")

    # Fallback to pdfplumber
    try:
        text = _parse_pdf_pdfplumber(file_path)
        return text
    except Exception as e:
        logger.error(f"PDF parsing failed: {e}")
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")


def _parse_pdf_pymupdf(file_path: str) -> str:
    """Extract text using PyMuPDF"""
    text_parts = []
    with fitz.open(file_path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def _parse_pdf_pdfplumber(file_path: str) -> str:
    """Extract text using pdfplumber"""
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def parse_docx(file_path: str) -> str:
    """
    Extract text from DOCX file using python-docx
    """
    try:
        doc = DocxDocument(file_path)
        text_parts = []

        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"DOCX parsing failed: {e}")
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")


def parse_txt(file_path: str) -> str:
    """
    Read text from TXT file with UTF-8 encoding
    """
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    except UnicodeDecodeError:
        # Try with different encoding if UTF-8 fails
        try:
            with open(file_path, 'r', encoding='latin-1') as f:
                return f.read()
        except Exception as e:
            logger.error(f"TXT parsing failed: {e}")
            raise ValueError(f"Failed to read text file: {str(e)}")
    except Exception as e:
        logger.error(f"TXT parsing failed: {e}")
        raise ValueError(f"Failed to read text file: {str(e)}")


def parse_document(file_path: str, file_type: str) -> str:
    """
    Parse document based on file type
    """
    if file_type == "pdf":
        return parse_pdf(file_path)
    elif file_type == "docx":
        return parse_docx(file_path)
    elif file_type == "txt":
        return parse_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")
